from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "00847E")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_word_example(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("智库来源导入示例（请替换示例内容）", level=1)
    document.add_heading("欧洲与国际战略", level=2)
    document.add_paragraph("示例智库甲")
    document.add_paragraph("https://example.org/publications/")
    document.add_paragraph("https://example.org/reports/")
    document.add_heading("表格和真实超链接示例", level=2)
    table = document.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("分类", "智库名称", "检索入口")):
        cell.text = value
    table.cell(1, 0).text = "国际战略"
    table.cell(1, 1).text = "示例智库乙"
    _add_hyperlink(
        table.cell(1, 2).paragraphs[0],
        "研究成果栏目",
        "https://research.example.net/papers/",
    )
    document.add_paragraph("提示：同一智库名称下可以连续列出多个 URL；导入后必须在预览页确认。")
    document.save(path)
    return path


def create_json_example(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema_version": 6,
        "prompt_version": "example",
        "configurations": [
            {
                "name": "示例配置（请替换）",
                "sources": [
                    {
                        "name": "示例智库（请替换）",
                        "group_name": "国际战略",
                        "homepage_url": "https://example.org/",
                        "enabled": True,
                        "entries": [
                            {
                                "url": "https://example.org/publications/",
                                "adapter_key": "auto",
                                "ordered_by_date": True,
                                "enabled": True,
                                "listing_selectors": ["article h2 a[href]"],
                                "container_selectors": ["article"],
                                "content_type_selectors": [".publication-type"],
                            },
                            {
                                "url": "https://example.org/reports/",
                                "adapter_key": "auto",
                                "ordered_by_date": False,
                                "enabled": True,
                            },
                        ],
                    }
                ],
            }
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
