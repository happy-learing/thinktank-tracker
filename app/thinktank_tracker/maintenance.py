from __future__ import annotations

import os
import shutil
import socket
import sqlite3
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

import httpx
from docx import Document
from openpyxl import Workbook
from sqlalchemy import select

from .config import get_api_key, load_config
from .crawler import bundled_chromium_executable
from .database import SessionLocal, engine
from .models import CrawlRun, Source, SourceEntry
from .paths import PATHS


def _chromium_network_probe(url: str) -> tuple[str, str]:
    browser_root = PATHS.root / "runtime" / "browsers"
    if browser_root.exists():
        os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(browser_root))
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            executable = bundled_chromium_executable(browser_root)
            browser = playwright.chromium.launch(
                headless=True,
                executable_path=str(executable) if executable else None,
            )
            page = browser.new_page(locale="en-US")
            response = page.goto(url, wait_until="domcontentloaded", timeout=30_000)
            status_code = response.status if response else 0
            browser.close()
        status = "pass" if 0 < status_code < 400 else "warn"
        return status, f"Chromium HTTP {status_code}: {url}"
    except Exception as exc:
        message = str(exc)
        if "ERR_NETWORK_ACCESS_DENIED" in message or "10013" in message:
            return "fail", f"本机拒绝应用进程联网: {url} ({message})"
        return "fail", f"Chromium 访问失败: {url} ({message})"


def _chromium_runtime_probe() -> tuple[str, str]:
    browser_root = PATHS.root / "runtime" / "browsers"
    if browser_root.exists():
        os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(browser_root))
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            executable = bundled_chromium_executable(browser_root)
            browser = playwright.chromium.launch(
                headless=True,
                executable_path=str(executable) if executable else None,
            )
            page = browser.new_page()
            page.set_content("<title>ThinkTankTracker Chromium Probe</title><main>ok</main>")
            title = page.title()
            browser.close()
        if title != "ThinkTankTracker Chromium Probe":
            return "fail", "Chromium 已启动，但页面渲染结果异常"
        return "pass", "Chromium 可启动并完成本地页面渲染"
    except Exception as exc:
        return "fail", f"Chromium 启动失败: {exc}"


@dataclass(frozen=True)
class DiagnosticItem:
    name: str
    status: str
    detail: str


def create_database_backup() -> Path:
    PATHS.ensure()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    sqlite_copy = PATHS.backups / f"tracker_{timestamp}.sqlite3"
    source = sqlite3.connect(PATHS.database)
    target = sqlite3.connect(sqlite_copy)
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()
    archive = PATHS.backups / f"智库追踪备份_{timestamp}.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as output:
        output.write(sqlite_copy, "tracker.sqlite3")
        config_file = PATHS.config / "app.json"
        if config_file.exists():
            output.write(config_file, "config/app.json")
    sqlite_copy.unlink(missing_ok=True)
    return archive


def purge_history_with_backup() -> tuple[Path, int]:
    """Remove run history while preserving source and model configuration."""
    archive = create_database_backup()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_backup = PATHS.backups / f"history_{timestamp}" / "exports"
    cache_index: list[str] = []
    for path in PATHS.cache.rglob("*"):
        if path.is_file():
            try:
                cache_index.append(f"{path.relative_to(PATHS.cache)}\t{path.stat().st_size}")
            except OSError:
                continue
    with zipfile.ZipFile(archive, "a", compression=zipfile.ZIP_DEFLATED) as output:
        output.writestr("cache-index.tsv", "\n".join(cache_index))
        for path in PATHS.exports.rglob("*"):
            if path.is_file():
                output.write(path, f"exports/{path.relative_to(PATHS.exports).as_posix()}")

    with SessionLocal.begin() as session:
        runs = session.scalars(select(CrawlRun)).all()
        removed = len(runs)
        for run in runs:
            session.delete(run)
        for source in session.scalars(select(Source)).all():
            source.last_successful_date = None
        for entry in session.scalars(select(SourceEntry)).all():
            entry.last_successful_date = None

    export_files = [path for path in PATHS.exports.rglob("*") if path.is_file()]
    if export_files:
        export_backup.mkdir(parents=True, exist_ok=True)
        for path in export_files:
            relative = path.relative_to(PATHS.exports)
            destination = export_backup / relative
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(path), str(destination))
    if PATHS.cache.exists():
        shutil.rmtree(PATHS.cache)
    PATHS.cache.mkdir(parents=True, exist_ok=True)
    PATHS.downloads.mkdir(parents=True, exist_ok=True)
    return archive, removed


def clean_expired_cache(days: int | None = None) -> int:
    retention = days if days is not None else load_config().raw_cache_days
    cutoff = datetime.now() - timedelta(days=retention)
    removed = 0
    for path in PATHS.downloads.rglob("*"):
        if path.is_file() and datetime.fromtimestamp(path.stat().st_mtime) < cutoff:
            path.unlink()
            removed += 1
    for directory in sorted((item for item in PATHS.downloads.rglob("*") if item.is_dir()), reverse=True):
        try:
            directory.rmdir()
        except OSError:
            pass
    return removed


def _port_available() -> bool:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1] > 0


def _export_smoke_test() -> str:
    with tempfile.TemporaryDirectory(dir=PATHS.workspace) as temp:
        directory = Path(temp)
        docx_path = directory / "test.docx"
        xlsx_path = directory / "test.xlsx"
        document = Document()
        document.add_paragraph("智库追踪导出测试")
        document.save(docx_path)
        workbook = Workbook()
        workbook.active.append(["智库追踪导出测试"])
        workbook.save(xlsx_path)
        if not docx_path.exists() or not xlsx_path.exists():
            raise RuntimeError("测试文件未生成")
    return "DOCX/XLSX 写入正常"


def _bundled_chromium_present(browser_root: Path) -> bool:
    return bundled_chromium_executable(browser_root) is not None


def run_diagnostics(sources: list[Source] | None = None, include_network: bool = False) -> list[DiagnosticItem]:
    PATHS.ensure()
    results: list[DiagnosticItem] = []
    try:
        probe = PATHS.workspace / ".write-test"
        probe.write_text("ok", encoding="ascii")
        probe.unlink()
        results.append(DiagnosticItem("工作目录", "pass", f"可写: {PATHS.workspace}"))
    except OSError as exc:
        results.append(DiagnosticItem("工作目录", "fail", str(exc)))
    try:
        with engine.connect() as connection:
            version = connection.exec_driver_sql("PRAGMA user_version").scalar_one()
        results.append(DiagnosticItem("SQLite", "pass", f"数据库版本 {version}，WAL 已启用"))
    except Exception as exc:
        results.append(DiagnosticItem("SQLite", "fail", str(exc)))
    results.append(DiagnosticItem("本地端口", "pass" if _port_available() else "fail", "可自动选择端口"))
    browser_root = PATHS.root / "runtime" / "browsers"
    browser_present = _bundled_chromium_present(browser_root)
    if browser_present:
        browser_status, browser_detail = _chromium_runtime_probe()
        results.append(DiagnosticItem("捆绑 Chromium", browser_status, browser_detail))
    else:
        results.append(
            DiagnosticItem(
                "捆绑 Chromium",
                "warn",
                "开发模式可使用 Playwright 已安装浏览器；发布包必须捆绑",
            )
        )
    try:
        results.append(DiagnosticItem("DOCX/XLSX 导出", "pass", _export_smoke_test()))
    except Exception as exc:
        results.append(DiagnosticItem("DOCX/XLSX 导出", "fail", str(exc)))
    windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    has_yahei = any(windows_fonts.glob("msyh*"))
    results.append(
        DiagnosticItem("中文字体", "pass" if has_yahei else "warn", "Microsoft YaHei" if has_yahei else "未发现微软雅黑，将由 Office 替代字体")
    )
    config = load_config()
    model_ready = bool(config.api_base_url and config.model_name and get_api_key())
    results.append(
        DiagnosticItem("模型 API", "pass" if model_ready else "warn", f"{config.model_name} / {config.api_base_url}" if model_ready else "尚未完整配置")
    )
    if include_network and sources:
        with httpx.Client(timeout=15, follow_redirects=True, headers={"User-Agent": "ThinkTankTracker/0.1"}) as client:
            for source in sources:
                entries = [entry for entry in source.entries if entry.enabled and entry.deleted_at is None]
                for index, entry in enumerate(entries, 1):
                    try:
                        response = client.get(entry.url)
                        status = "pass" if response.status_code < 400 else "warn"
                        detail = f"HTTP {response.status_code}: {entry.url}"
                    except Exception as exc:
                        status, detail = "fail", f"{entry.url}: {exc}"
                    label = f"连通性: {source.name}" if len(entries) == 1 else f"连通性: {source.name} #{index}"
                    results.append(DiagnosticItem(f"HTTP {label}", status, detail))
                    browser_status, browser_detail = _chromium_network_probe(entry.url)
                    results.append(DiagnosticItem(f"Chromium {label}", browser_status, browser_detail))
    if include_network and model_ready:
        endpoint = config.api_base_url.rstrip("/")
        if endpoint.endswith("/chat/completions"):
            endpoint = endpoint[: -len("/chat/completions")]
        elif endpoint.endswith("/responses"):
            endpoint = endpoint[: -len("/responses")]
        try:
            response = httpx.get(
                f"{endpoint}/models",
                headers={"Authorization": f"Bearer {get_api_key()}"},
                timeout=20,
            )
            status = "pass" if response.status_code < 400 else "warn"
            detail = f"HTTP {response.status_code} /models"
        except Exception as exc:
            status, detail = "fail", str(exc)
        results.append(DiagnosticItem("模型 API 连通性", status, detail))
    return results
