from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from .constants import CATEGORY_LABELS, MICRO_CATEGORY_ORDER, CandidateStatus, CoverageStatus, DateScope
from .models import (
    Candidate,
    CandidateAttachment,
    CandidateDiscovery,
    CrawlRun,
    DateAssistBatch,
    ListingPageAudit,
    ModelSuggestion,
    Source,
    SourceEntryRun,
    SourceRun,
)
from .paths import PATHS
from .selection import fact_issues, run_policy


NAVY = "1F4E79"
TEAL = "0F6B78"
LIGHT_BLUE = "D9EAF3"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "6B7280"
WHITE = "FFFFFF"
RED = "B42318"
REPORT_TEMPLATE_PATH = Path(__file__).with_name("assets") / "report_template.docx"


@dataclass(frozen=True)
class ReportArticle:
    authors: tuple[str, ...]
    title_en: str
    title_zh: str
    published_date: date
    url: str


def _set_run_font(run, latin: str = "Calibri", east_asia: str = "Microsoft YaHei", size: int = 10) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    _set_run_font(run, size=9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    tail = paragraph.add_run(" 页")
    _set_run_font(tail, size=9)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    tail.font.color.rgb = RGBColor.from_string(MID_GRAY)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0000FF")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "22")
    fonts = OxmlElement("w:rFonts")
    for key, value in (
        ("ascii", "Times New Roman"),
        ("hAnsi", "Times New Roman"),
        ("cs", "Times New Roman"),
        ("eastAsia", "宋体"),
    ):
        fonts.set(qn(f"w:{key}"), value)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "21")
    properties.extend([style, fonts, color, underline, size, size_cs])
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.1)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color in (("Title", 28, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 13, TEAL)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    if "Article Title" not in [style.name for style in styles]:
        article_style = styles.add_style("Article Title", WD_STYLE_TYPE.PARAGRAPH)
        article_style.font.name = "Calibri"
        article_style.font.size = Pt(11.5)
        article_style.font.bold = True
        article_style.font.color.rgb = RGBColor.from_string(NAVY)
        article_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        article_style.paragraph_format.keep_with_next = True
        article_style.paragraph_format.space_before = Pt(9)
        article_style.paragraph_format.space_after = Pt(3)
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.8))
    table.columns[0].width = Inches(4.8)
    table.columns[1].width = Inches(2.0)
    left = table.cell(0, 0).paragraphs[0]
    left_run = left.add_run("THINK TANK RESEARCH TRACKER")
    _set_run_font(left_run, size=8)
    left_run.bold = True
    left_run.font.color.rgb = RGBColor.from_string(NAVY)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_run = right.add_run("选题参考")
    _set_run_font(right_run, size=8)
    right_run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    _add_page_number(section.footer.paragraphs[0])


def _article_title(candidate: Candidate) -> str:
    review = candidate.review
    english = (review.title_en or candidate.title_en or candidate.title_original).strip()
    chinese = (review.title_zh or candidate.title_zh).strip()
    if english and chinese and english.casefold() != chinese.casefold():
        return f"{english}\n{chinese}"
    return english or chinese or "（标题待补）"


def _json_string_list(value: str) -> list[str]:
    try:
        parsed = json.loads(value or "[]")
    except (json.JSONDecodeError, TypeError):
        return []
    if not isinstance(parsed, list):
        return []
    return [str(item).strip() for item in parsed if str(item).strip()]


def _set_reference_font(run, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _english_date(value: date) -> str:
    return f"{value.strftime('%B')} {value.day}, {value.year}"


def _chinese_date(value: date) -> str:
    return f"{value.year}年{value.month}月{value.day}日"


def _report_candidate_errors(candidates: list[Candidate], run: CrawlRun) -> dict[int, list[str]]:
    failures: dict[int, list[str]] = {}
    for candidate in candidates:
        missing: list[str] = []
        review = candidate.review
        if review is None:
            missing.append("人工审核记录")
        else:
            if not _json_string_list(review.authors_json):
                missing.append("作者")
            if not review.title_en.strip():
                missing.append("英文标题")
            if not review.title_zh.strip():
                missing.append("中文标题")
        if candidate.date_status != "confirmed" or candidate.published_date is None:
            missing.append("确认发布日期")
        elif not run.start_date <= candidate.published_date <= run.end_date:
            missing.append("任务日期范围内的发布日期")
        if missing:
            failures[candidate.id] = missing
    return failures


def build_weekly_report_document(
    start_date: date,
    end_date: date,
    groups: list[tuple[str, list[ReportArticle]]],
) -> Document:
    if not REPORT_TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"内置 Word 模板不存在：{REPORT_TEMPLATE_PATH}")
    document = Document(REPORT_TEMPLATE_PATH)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.25
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title_run = title.add_run(
        f"智库周览（{_chinese_date(start_date)}-{_chinese_date(end_date)}）"
    )
    _set_reference_font(title_run, bold=True)

    for source_name, articles in groups:
        heading = document.add_paragraph(style="Heading 1")
        heading.paragraph_format.line_spacing = 1.25
        heading.paragraph_format.keep_with_next = True
        heading_run = heading.add_run(source_name)
        _set_reference_font(heading_run, bold=True)

        numbered = len(articles) > 1
        for index, article in enumerate(articles, 1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = True
            if numbered:
                _set_reference_font(paragraph.add_run(f"{index}. "))
            _set_reference_font(paragraph.add_run(f"{', '.join(article.authors)}, "))
            title_text = f"{article.title_en}（{article.title_zh}）"
            _set_reference_font(paragraph.add_run(title_text), italic=True)
            _set_reference_font(
                paragraph.add_run(f", {_english_date(article.published_date)}.")
            )

            url_paragraph = document.add_paragraph()
            url_paragraph.paragraph_format.line_spacing = 1.25
            url_paragraph.paragraph_format.space_before = Pt(0)
            url_paragraph.paragraph_format.space_after = Pt(0)
            url_paragraph.paragraph_format.keep_together = True
            _add_hyperlink(url_paragraph, article.url, article.url)
    return document


def _entry_snapshot(entry_run: SourceEntryRun) -> dict[str, object]:
    try:
        value = json.loads(entry_run.entry_config_snapshot_json or "{}")
    except json.JSONDecodeError:
        return {}
    return value if isinstance(value, dict) else {}


def _attachment_audit_cells(value: str) -> list[object]:
    try:
        audit = json.loads(value or "{}")
    except (json.JSONDecodeError, TypeError):
        audit = {}
    if not isinstance(audit, dict):
        audit = {}
    engine = str(audit.get("final_engine") or "")
    engine_label = {
        "anydoc": "AnyDoc",
        "pypdf": "pypdf",
        "pypdf_legacy": "pypdf（历史）",
        "none": "未生成正文",
    }.get(engine, engine or "未记录")
    attempts = audit.get("attempts")
    if engine == "none" and isinstance(attempts, list) and attempts:
        last_attempt = attempts[-1]
        if isinstance(last_attempt, dict) and last_attempt.get("engine"):
            engine_label = f"{last_attempt['engine']}（未生成正文）"
    duration = audit.get("duration_ms")
    if not isinstance(duration, (int, float)):
        duration = ""
    ocr_pages = audit.get("ocr_pages")
    ocr_text = ", ".join(str(item) for item in ocr_pages) if isinstance(ocr_pages, list) else ""
    fallback = str(audit.get("fallback_detail") or audit.get("fallback_reason") or "")
    return [
        engine_label,
        str(audit.get("final_engine_version") or ""),
        duration,
        fallback,
        ocr_text,
        str(audit.get("final_source") or ""),
    ]


def export_report(session: Session, run_id: int) -> Path:
    run = session.get(CrawlRun, run_id)
    if not run or run.deleted_at is not None:
        raise ValueError("检索任务不存在")
    candidates = list(session.scalars(
        select(Candidate)
        .where(
            Candidate.run_id == run_id,
            Candidate.status == CandidateStatus.APPROVED,
        )
        .options(selectinload(Candidate.source), selectinload(Candidate.review))
        .order_by(Candidate.id)
    ).all())
    failures = _report_candidate_errors(candidates, run)
    if failures:
        detail = "；".join(
            f"候选 {candidate_id}: {', '.join(fields)}"
            for candidate_id, fields in failures.items()
        )
        raise ValueError(f"正式 Word 导出失败，已批准条目不完整：{detail}")
    source_runs = list(session.scalars(
        select(SourceRun)
        .where(SourceRun.run_id == run_id)
        .options(selectinload(SourceRun.source))
        .order_by(SourceRun.id)
    ).all())
    source_order = {row.source_id: index for index, row in enumerate(source_runs)}
    source_names = {
        row.source_id: (row.source_name_snapshot or row.source.name)
        for row in source_runs
    }
    grouped: dict[int, list[Candidate]] = {}
    for candidate in candidates:
        grouped.setdefault(candidate.source_id, []).append(candidate)
    ordered_source_ids = sorted(
        grouped,
        key=lambda source_id: (source_order.get(source_id, len(source_order)), source_id),
    )

    report_groups: list[tuple[str, list[ReportArticle]]] = []
    for source_id in ordered_source_ids:
        rows = sorted(
            grouped[source_id],
            key=lambda item: (
                item.published_date,
                item.review.title_en.casefold(),
                item.id,
            ),
        )
        report_groups.append(
            (
                source_names.get(source_id, rows[0].source.name),
                [
                    ReportArticle(
                        authors=tuple(_json_string_list(candidate.review.authors_json)),
                        title_en=candidate.review.title_en,
                        title_zh=candidate.review.title_zh,
                        published_date=candidate.published_date,
                        url=candidate.original_url,
                    )
                    for candidate in rows
                ],
            )
        )

    document = build_weekly_report_document(
        run.start_date,
        run.end_date,
        report_groups,
    )

    PATHS.exports.mkdir(parents=True, exist_ok=True)
    path = PATHS.exports / (
        f"{_chinese_date(run.start_date)}-{_chinese_date(run.end_date)}.docx"
    )
    document.save(path)
    return path


def draft_issues(candidate: Candidate, run: CrawlRun) -> list[str]:
    issues = fact_issues(candidate, run.start_date, run.end_date, run_policy(run))
    if not _json_string_list(candidate.authors_json):
        issues.append("作者待核实")
    if not candidate.title_en.strip() or not candidate.title_zh.strip():
        issues.append("双语标题待补齐")
    try:
        audit = json.loads(candidate.metadata_json or "{}").get("selection", {})
    except (ValueError, AttributeError):
        audit = {}
    if not (audit.get("evidence_verified") and audit.get("input_complete")
            and audit.get("verification_status") == "supported"
            and audit.get("research_substance") == "substantive"):
        issues.append("全文分析及证据复核未通过")
    if not candidate.assessment or candidate.assessment.decision != "include":
        issues.append("尚无纳入结论")
    if candidate.review is not None:
        issues.append("存在人工处理记录，草稿遵循人工结果")
    return list(dict.fromkeys(issues))


def export_draft_report(session: Session, run_id: int) -> Path:
    run = session.get(CrawlRun, run_id)
    if not run or run.deleted_at is not None:
        raise ValueError("检索任务不存在")
    rows = list(session.scalars(select(Candidate).where(Candidate.run_id == run_id)
        .options(selectinload(Candidate.source), selectinload(Candidate.review), selectinload(Candidate.assessment))
        .order_by(Candidate.id)).all())
    sources = list(session.scalars(select(SourceRun).where(SourceRun.run_id == run_id)
        .options(selectinload(SourceRun.source)).order_by(SourceRun.id)).all())
    names = {r.source_id: r.source_name_snapshot or r.source.name for r in sources}
    order = {r.source_id: i for i, r in enumerate(sources)}
    groups = {}
    exceptions = []
    approved = automatic = 0
    for row in rows:
        if row.status in {CandidateStatus.REJECTED, CandidateStatus.HARD_EXCLUDED, CandidateStatus.MODEL_EXCLUDED}:
            continue
        if row.status == CandidateStatus.APPROVED:
            problems = _report_candidate_errors([row], run).get(row.id, [])
            fields = row.review
        else:
            problems = draft_issues(row, run)
            fields = row
        if problems:
            exceptions.append((row, problems))
            continue
        groups.setdefault(row.source_id, []).append(ReportArticle(
            authors=tuple(_json_string_list(fields.authors_json)), title_en=fields.title_en,
            title_zh=fields.title_zh, published_date=row.published_date, url=row.original_url))
        approved += row.status == CandidateStatus.APPROVED
        automatic += row.status != CandidateStatus.APPROVED
    doc = build_weekly_report_document(run.start_date, run.end_date, [
        (names.get(source_id, "来源待核实"), sorted(articles,key=lambda a:(a.published_date,a.title_en)))
        for source_id, articles in sorted(groups.items(),key=lambda item:(order.get(item[0],len(order)),item[0]))])
    title = doc.paragraphs[0]
    title.add_run(" — 自动草稿")
    note = title.insert_paragraph_before(
        f"任务 #{run.id} · 自动筛选 {automatic} 条 · 人工批准 {approved} 条 · 待处理 {len(exceptions)} 条。"
        "自动条目通过程序核验与模型复核，尚未经人工逐篇批准。")
    _set_run_font(note.runs[0], size=9)
    coverage = "；".join(f"{r.source_name_snapshot or r.source.name}: {r.status}（{r.stop_reason or '尚未结束'}）" for r in sources)
    cover_note = title.insert_paragraph_before("来源覆盖：" + (coverage or "未记录，不能确认完整性"))
    _set_run_font(cover_note.runs[0], size=9)
    doc.core_properties.title = f"任务 {run.id} 自动草稿"
    if not automatic and not approved:
        doc.add_paragraph("当前没有满足导出条件的条目。请查看下方待处理内容和来源覆盖，不能据此断言没有符合要求的研究。")
    if exceptions:
        doc.add_heading("待处理条目", level=1)
        for row, problems in exceptions:
            doc.add_paragraph(f"候选 {row.id}：{row.title_original}\n" + "；".join(problems))
            _add_hyperlink(doc.add_paragraph(), row.original_url, row.original_url)
    PATHS.exports.mkdir(parents=True,exist_ok=True)
    path = PATHS.exports / f"草稿-任务{run.id}-{datetime.now():%Y%m%d-%H%M%S-%f}.docx"
    doc.save(path)
    return path


def _sheet(ws, headers: list[str], rows: list[list[object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    thin = Side(style="thin", color="D0D5DD")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 28
    for index, header in enumerate(headers, 1):
        values = [str(ws.cell(row=row, column=index).value or "") for row in range(1, min(ws.max_row, 100) + 1)]
        width = min(max(len(header) + 2, max((len(value) for value in values), default=0) + 2), 55)
        ws.column_dimensions[get_column_letter(index)].width = max(10, width)


def export_audit_workbook(session: Session, run_id: int) -> Path:
    run = session.get(CrawlRun, run_id)
    if not run or run.deleted_at is not None:
        raise ValueError("检索任务不存在")
    source_runs = session.scalars(
        select(SourceRun)
        .where(SourceRun.run_id == run_id)
        .options(
            selectinload(SourceRun.source),
            selectinload(SourceRun.entry_runs).selectinload(SourceEntryRun.source_entry),
            selectinload(SourceRun.entry_runs).selectinload(SourceEntryRun.page_audits),
        )
        .order_by(SourceRun.id)
    ).all()
    candidates = session.scalars(
        select(Candidate)
        .where(Candidate.run_id == run_id)
        .options(
            selectinload(Candidate.source),
            selectinload(Candidate.assessment),
            selectinload(Candidate.review),
            selectinload(Candidate.discoveries).selectinload(CandidateDiscovery.source_entry),
            selectinload(Candidate.attachments),
            selectinload(Candidate.model_suggestions),
        )
        .order_by(Candidate.source_id, Candidate.id)
    ).all()
    source_names = {
        row.source_id: (row.source_name_snapshot or row.source.name)
        for row in source_runs
    }
    entry_urls = {
        entry_run.source_entry_id: (entry_run.entry_url_snapshot or entry_run.source_entry.url)
        for row in source_runs
        for entry_run in row.entry_runs
    }
    wb = Workbook()
    overview = wb.active
    overview.title = "概览"
    summary_rows = [
        ["检索配置", run.configuration_name_snapshot],
        ["日期范围", f"{run.start_date.isoformat()} 至 {run.end_date.isoformat()}"],
        ["任务状态", run.status],
        ["来源数量", len(source_runs)],
        ["检索入口数量", sum(len(item.entry_runs) for item in source_runs)],
        ["候选总数", len(candidates)],
        ["日期范围内全部内容", sum(item.date_scope == DateScope.IN_WINDOW for item in candidates)],
        ["日期待确认", sum(item.date_scope == DateScope.DATE_UNKNOWN for item in candidates)],
        ["日期范围外", sum(item.date_scope in {DateScope.AFTER_END, DateScope.BEFORE_START} for item in candidates)],
        ["人工批准", sum(item.status == CandidateStatus.APPROVED for item in candidates)],
        ["待人工处理", sum(item.status == CandidateStatus.NEEDS_REVIEW for item in candidates)],
        ["覆盖不完整/失败", sum(item.status != CoverageStatus.COMPLETE for item in source_runs)],
        ["生成时间", datetime.now().isoformat(timespec="seconds")],
    ]
    _sheet(overview, ["指标", "值"], summary_rows)
    coverage_rows = [
        [
            row.source_name_snapshot or row.source.name,
            row.status,
            len(row.entry_runs),
            row.pages_scanned,
            row.links_discovered,
            row.earliest_date.isoformat() if row.earliest_date else "",
            row.latest_date.isoformat() if row.latest_date else "",
            row.stop_reason,
            row.error_message,
        ]
        for row in source_runs
    ]
    _sheet(
        wb.create_sheet("智库覆盖"),
        ["智库", "覆盖状态", "入口数量", "访问页数", "发现链接", "最早日期", "最晚日期", "停止依据", "错误"],
        coverage_rows,
    )
    entry_coverage_rows = [
        [
            row.source_name_snapshot or row.source.name,
            entry_run.entry_url_snapshot or entry_run.source_entry.url,
            _entry_snapshot(entry_run).get("adapter_key", entry_run.source_entry.adapter_key),
            entry_run.status,
            entry_run.pages_scanned,
            entry_run.links_discovered,
            entry_run.earliest_date.isoformat() if entry_run.earliest_date else "",
            entry_run.latest_date.isoformat() if entry_run.latest_date else "",
            entry_run.ordering_status,
            entry_run.boundary_page_number or "",
            entry_run.stop_reason,
            entry_run.error_message,
        ]
        for row in source_runs
        for entry_run in row.entry_runs
    ]
    _sheet(
        wb.create_sheet("入口覆盖"),
        ["智库", "检索入口", "适配器", "覆盖状态", "访问页数", "发现链接", "最早日期", "最晚日期", "顺序验证", "边界页", "停止依据", "错误"],
        entry_coverage_rows,
    )
    page_rows = [
        [
            row.source_name_snapshot or row.source.name,
            entry_run.entry_url_snapshot or entry_run.source_entry.url,
            audit.page_number,
            audit.page_url,
            audit.cards_found,
            audit.accepted_count,
            audit.rejected_count,
            audit.in_window_count,
            audit.after_end_count,
            audit.before_start_count,
            audit.date_unknown_count,
            audit.earliest_date.isoformat() if audit.earliest_date else "",
            audit.latest_date.isoformat() if audit.latest_date else "",
            audit.ordering_status,
            "是" if audit.is_boundary_page else "否",
            audit.stop_evidence,
            audit.sample_links_json,
        ]
        for row in source_runs
        for entry_run in row.entry_runs
        for audit in entry_run.page_audits
    ]
    _sheet(
        wb.create_sheet("逐页发现审计"),
        ["智库", "入口", "页码", "栏目页", "卡片数", "接收数", "拒绝数", "范围内", "过新", "过旧", "日期不明", "最早日期", "最晚日期", "顺序", "边界页", "停止证据", "样例链接"],
        page_rows,
    )
    no_match_rows = []
    for row in source_runs:
        approved = sum(item.source_id == row.source_id and item.status == CandidateStatus.APPROVED for item in candidates)
        if not approved:
            statement = (
                f"该智库在{run.start_date:%Y.%m.%d}-{run.end_date:%m.%d}期间暂无符合条件的深度研究"
                if row.status == CoverageStatus.COMPLETE
                else "覆盖不完整，不能判定暂无符合内容"
            )
            no_match_rows.append(
                [
                    row.source_name_snapshot or row.source.name,
                    row.status,
                    statement,
                    row.error_message,
                ]
            )
    _sheet(wb.create_sheet("无符合项智库"), ["智库", "覆盖状态", "结论", "说明"], no_match_rows)
    date_rows = [
        [
            source_names.get(item.source_id, item.source.name),
            item.title_original,
            item.date_status,
            item.date_evidence,
            item.original_url,
        ]
        for item in candidates
        if item.date_status != "confirmed"
    ]
    _sheet(wb.create_sheet("日期待确认"), ["智库", "标题", "日期状态", "日期证据", "原文链接"], date_rows)
    failure_rows = [
        [
            source_names.get(item.source_id, item.source.name),
            item.title_original,
            item.extraction_status,
            item.hard_filter_reason,
            item.fetch_audit_json,
            item.original_url,
        ]
        for item in candidates
        if item.extraction_status not in {"complete"}
    ]
    _sheet(
        wb.create_sheet("抓取失败"),
        ["智库", "标题", "提取状态", "问题", "URL 请求尝试", "原文链接"],
        failure_rows,
    )
    candidate_rows = []
    for item in candidates:
        candidate_rows.append(
            [
                item.id,
                source_names.get(item.source_id, item.source.name),
                item.title_original,
                "\n".join(_json_string_list(item.authors_json)),
                item.author_evidence,
                "\n".join(_json_string_list(item.review.authors_json)) if item.review else "",
                item.published_date.isoformat() if item.published_date else "",
                item.date_scope,
                item.date_status,
                item.content_type,
                item.content_type_evidence,
                "\n".join(
                    row.listing_content_type_evidence
                    for row in item.discoveries
                    if row.listing_content_type_evidence
                ),
                item.content_kind_status,
                item.language,
                item.english_word_count,
                item.chinese_char_count,
                item.length_status,
                item.extraction_status,
                item.topic_status,
                item.review_status,
                item.status,
                item.hard_filter_reason,
                item.assessment.decision if item.assessment else "",
                item.assessment.primary_category if item.assessment else "",
                item.review.decision if item.review else "",
                item.review.notes if item.review else "",
                item.content_hash,
                "\n".join(
                    dict.fromkeys(
                        entry_urls.get(row.source_entry_id, row.source_entry.url)
                        for row in item.discoveries
                    )
                ),
                "\n".join(dict.fromkeys(row.listing_url for row in item.discoveries if row.listing_url)),
                item.fetch_audit_json,
                item.original_url,
            ]
        )
    _sheet(
        wb.create_sheet("全部正式内容"),
        [
            "ID", "智库", "标题", "自动作者", "作者证据", "人工作者", "发布日期", "日期范围", "日期状态", "体裁", "体裁证据", "列表体裁证据", "体裁状态", "语言", "英文词数", "中文字符数",
            "篇幅状态", "提取状态", "主题状态", "审核状态", "旧版兼容状态", "筛选/异常原因", "模型决定", "模型主分类", "人工决定", "人工备注", "正文哈希",
            "发现入口", "发现栏目页", "URL 请求尝试", "原文链接",
        ],
        candidate_rows,
    )
    full_headers = [
        "ID", "智库", "标题", "自动作者", "作者证据", "人工作者", "发布日期", "日期范围", "日期状态", "体裁", "体裁证据", "列表体裁证据", "体裁状态", "语言", "英文词数", "中文字符数",
        "篇幅状态", "提取状态", "主题状态", "审核状态", "旧版兼容状态", "筛选/异常原因", "模型决定", "模型主分类", "人工决定", "人工备注", "正文哈希",
        "发现入口", "发现栏目页", "URL 请求尝试", "原文链接",
    ]
    _sheet(
        wb.create_sheet("日期范围内完整清单"),
        full_headers,
        [row for row, item in zip(candidate_rows, candidates) if item.date_scope == DateScope.IN_WINDOW],
    )
    _sheet(
        wb.create_sheet("范围外内容"),
        full_headers,
        [row for row, item in zip(candidate_rows, candidates) if item.date_scope in {DateScope.AFTER_END, DateScope.BEFORE_START}],
    )
    model_rows = [
        [
            source_names.get(item.source_id, item.source.name),
            item.title_original,
            item.assessment.provider_model,
            item.assessment.input_chars,
            item.assessment.confidence,
            item.assessment.error_message,
        ]
        for item in candidates
        if item.assessment
    ]
    _sheet(wb.create_sheet("模型调用统计"), ["智库", "标题", "模型", "输入字符", "置信度", "错误"], model_rows)
    attachment_rows = [
        [
            source_names.get(item.source_id, item.source.name),
            item.title_original,
            attachment.role,
            attachment.link_text,
            attachment.status_code or "",
            attachment.content_type,
            attachment.extraction_status,
            *_attachment_audit_cells(attachment.extraction_audit_json),
            "是" if attachment.used_as_body else "否",
            attachment.content_hash,
            attachment.error_message,
            attachment.url,
        ]
        for item in candidates
        for attachment in item.attachments
    ]
    _sheet(
        wb.create_sheet("PDF附件"),
        [
            "智库", "文章", "附件角色", "链接文字", "HTTP状态", "响应类型", "提取状态",
            "解析器", "解析器版本", "解析耗时（毫秒）", "回退原因", "OCR页码", "最终正文来源",
            "作为正文", "哈希", "错误", "附件URL",
        ],
        attachment_rows,
    )
    suggestion_rows = [
        [
            source_names.get(item.source_id, item.source.name),
            item.title_original,
            suggestion.provider_model,
            suggestion.created_at.isoformat(timespec="seconds"),
            suggestion.decision,
            suggestion.suggested_date.isoformat() if suggestion.suggested_date else "",
            suggestion.date_evidence,
            suggestion.content_type,
            suggestion.primary_category,
            suggestion.title_zh,
            suggestion.core_claim_zh,
            suggestion.confidence,
            suggestion.input_chars,
            suggestion.error_message,
        ]
        for item in candidates
        for suggestion in item.model_suggestions
    ]
    _sheet(
        wb.create_sheet("人工触发模型建议"),
        ["智库", "文章", "模型", "调用时间", "建议决定", "建议日期", "日期证据", "建议体裁", "主分类", "中文标题", "核心观点", "置信度", "输入字符", "错误"],
        suggestion_rows,
    )
    date_assists = session.scalars(
        select(DateAssistBatch).where(DateAssistBatch.run_id == run_id).order_by(DateAssistBatch.id)
    ).all()
    _sheet(
        wb.create_sheet("模型日期辅助"),
        ["入口ID", "栏目页", "内容URL", "模型", "建议日期", "原文证据", "证据已验证", "错误", "调用时间"],
        [
            [
                item.source_entry_id,
                item.listing_url,
                item.item_url,
                item.provider_model,
                item.suggested_date.isoformat() if item.suggested_date else "",
                item.evidence_quote,
                "是" if item.verified else "否",
                item.error_message,
                item.created_at.isoformat(timespec="seconds"),
            ]
            for item in date_assists
        ],
    )
    selection_rows = []
    for item in candidates:
        try:
            audit = json.loads(item.metadata_json or "{}").get("selection", {})
        except (ValueError, AttributeError):
            audit = {}
        selection_rows.append([
            item.id, item.title_original, audit.get("policy", {}).get("mode", "历史未记录"),
            audit.get("research_substance", ""), audit.get("input_complete", False),
            audit.get("evidence_verified", False), audit.get("verification_status", "未记录"),
            audit.get("verification_reason", ""), audit.get("model_calls", ""),
            audit.get("input_chars", ""), audit.get("input_sha256", ""),
            "；".join(audit.get("fact_issues", [])),
        ])
    _sheet(wb.create_sheet("研究判定核验"),
        ["候选ID", "文章", "任务模式", "实质分析", "全文覆盖", "引文回查", "复核结果",
         "复核理由", "逻辑调用次数", "最终输入字符", "最终输入哈希", "未满足条件"], selection_rows)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    PATHS.exports.mkdir(parents=True, exist_ok=True)
    path = PATHS.exports / f"智库检索审计_{run.start_date.isoformat()}_{run.end_date.isoformat()}.xlsx"
    wb.save(path)
    return path
