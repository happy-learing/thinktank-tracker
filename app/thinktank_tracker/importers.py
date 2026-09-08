from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Iterator
from urllib.parse import urlsplit

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .schemas import ImportedSource, ImportedSourceEntry
from .source_catalog import known_source_for
from .url_utils import canonicalize_url, extract_urls, normalize_fetch_url


NUMBER_PREFIX_RE = re.compile(r"^\s*(?:[（(]?\d+[）).、．]\s*|[-–—•]\s*)")


@dataclass
class Block:
    text: str
    urls: list[str]
    style_name: str = ""
    cells: list["Block"] = field(default_factory=list)
    table_headers: list[str] = field(default_factory=list)


def _iter_blocks(parent: DocumentObject | _Cell) -> Iterator[Paragraph | Table]:
    parent_element = parent.element.body if isinstance(parent, DocumentObject) else parent._tc
    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _paragraph_text_and_urls(paragraph: Paragraph) -> tuple[str, list[str]]:
    text = "".join(node.text or "" for node in paragraph._p.xpath(".//w:t"))
    urls = extract_urls(text)
    for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
        rel_id = hyperlink.get(qn("r:id"))
        if not rel_id:
            continue
        relation = paragraph.part.rels.get(rel_id)
        if relation and relation.target_ref.startswith(("http://", "https://")):
            urls.append(relation.target_ref)
    return text.strip(), list(dict.fromkeys(urls))


def _document_blocks(document: DocumentObject) -> Iterator[Block]:
    for item in _iter_blocks(document):
        if isinstance(item, Paragraph):
            text, urls = _paragraph_text_and_urls(item)
            style = item.style.name if item.style is not None else ""
            yield Block(text=text, urls=urls, style_name=style)
            continue
        table_headers: list[str] = []
        for row_index, row in enumerate(item.rows):
            cell_blocks: list[Block] = []
            for cell in row.cells:
                text_parts: list[str] = []
                urls: list[str] = []
                for paragraph in cell.paragraphs:
                    text, paragraph_urls = _paragraph_text_and_urls(paragraph)
                    if text:
                        text_parts.append(text)
                    urls.extend(paragraph_urls)
                cell_blocks.append(
                    Block(
                        text=" ".join(text_parts).strip(),
                        urls=list(dict.fromkeys(urls)),
                        style_name="table-cell",
                    )
                )
            row_urls = [url for block in cell_blocks for url in block.urls]
            row_text = " | ".join(block.text for block in cell_blocks if block.text)
            header_values = [block.text.strip().lower() for block in cell_blocks]
            header_terms = {
                "智库", "智库名称", "名称", "分类", "类别", "检索入口", "入口", "栏目",
                "链接", "网址", "url", "source", "category", "group", "think tank", "name",
            }
            is_header = row_index == 0 and any(value in header_terms for value in header_values)
            if is_header:
                table_headers = header_values
                yield Block(
                    text=row_text,
                    urls=[],
                    style_name="table-header",
                    cells=cell_blocks,
                    table_headers=table_headers,
                )
                continue
            yield Block(
                text=row_text,
                urls=list(dict.fromkeys(row_urls)),
                style_name="table-row",
                cells=cell_blocks,
                table_headers=table_headers,
            )


def _clean_name(text: str, urls: list[str]) -> str:
    result = text
    for url in urls:
        result = result.replace(url, " ")
    result = NUMBER_PREFIX_RE.sub("", result)
    result = result.replace("|", " ").strip(" \t:-–—")
    return re.sub(r"\s+", " ", result)


def _fallback_name(url: str) -> str:
    hostname = (urlsplit(url).hostname or url).removeprefix("www.")
    return hostname


def _group_key(name: str, url: str) -> str:
    known = known_source_for(url)
    if known:
        return f"known:{known.institution_key}"
    normalized_name = re.sub(r"\s+", " ", name).strip().casefold()
    hostname = (urlsplit(url).hostname or "").lower().removeprefix("www.")
    return f"custom:{hostname}:{normalized_name}"


def parse_docx_sources(source: str | Path | bytes | BinaryIO) -> list[ImportedSource]:
    if isinstance(source, bytes):
        document = Document(io.BytesIO(source))
    else:
        document = Document(source)

    current_group = ""
    active_name = ""
    imported: dict[str, ImportedSource] = {}
    seen_entries: set[tuple[str, str]] = set()
    seen_urls: set[str] = set()

    for block in _document_blocks(document):
        text = block.text.strip()
        urls = block.urls or extract_urls(text)
        if block.style_name == "table-header":
            active_name = ""
            continue
        is_heading = block.style_name.lower().startswith("heading") or block.style_name.startswith("标题")
        if is_heading and text and not urls:
            current_group = _clean_name(text, [])
            active_name = ""
            continue
        if not urls:
            if text:
                active_name = _clean_name(text, [])
            continue

        block_group = current_group
        table_name = ""
        if block.style_name == "table-row" and block.cells:
            headers = block.table_headers

            def header_index(keywords: tuple[str, ...]) -> int | None:
                return next(
                    (index for index, value in enumerate(headers) if any(key in value for key in keywords)),
                    None,
                )

            name_index = header_index(("智库", "名称", "think tank", "institute", "source"))
            group_index = header_index(("分类", "类别", "category", "group"))
            if name_index is not None and name_index < len(block.cells):
                table_name = _clean_name(block.cells[name_index].text, block.cells[name_index].urls)
            if group_index is not None and group_index < len(block.cells):
                block_group = _clean_name(block.cells[group_index].text, block.cells[group_index].urls) or current_group
            if not table_name:
                url_indexes = [index for index, cell in enumerate(block.cells) if cell.urls]
                preceding = [
                    block.cells[index].text
                    for index in range((url_indexes[0] if url_indexes else len(block.cells)) - 1, -1, -1)
                    if block.cells[index].text and not block.cells[index].urls
                ]
                table_name = _clean_name(preceding[0], []) if preceding else ""

        inline_name = table_name or _clean_name(text, urls)
        name = inline_name or active_name
        if inline_name:
            active_name = inline_name
        for url in urls:
            fetch_url = normalize_fetch_url(url)
            resolved_name = name or _fallback_name(fetch_url)
            key = _group_key(resolved_name, fetch_url)
            canonical = canonicalize_url(fetch_url)
            if canonical in seen_urls or (key, canonical) in seen_entries:
                continue
            seen_urls.add(canonical)
            seen_entries.add((key, canonical))
            known = known_source_for(fetch_url)
            entry = ImportedSourceEntry(
                url=fetch_url,
                adapter_key=known.adapter_key if known else "generic",
                ordered_by_date=known.ordered_by_date if known else True,
            )
            if key not in imported:
                imported[key] = ImportedSource(
                    name=resolved_name,
                    group_name=block_group,
                    homepage_url=known.homepage_url if known else fetch_url,
                    entries=[entry],
                )
            else:
                imported[key].entries.append(entry)
            imported[key].publication_url = imported[key].entries[0].url
            imported[key].adapter_key = imported[key].entries[0].adapter_key
            imported[key].ordered_by_date = imported[key].entries[0].ordered_by_date

    return list(imported.values())
