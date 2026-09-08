from __future__ import annotations

import io
from pathlib import Path

from docx import Document

from thinktank_tracker.importers import parse_docx_sources
from thinktank_tracker.example_files import create_word_example
from thinktank_tracker.source_service import parse_sources_json


def test_parses_current_sample() -> None:
    sample = Path(__file__).resolve().parents[2] / "追踪智库.docx"
    sources = parse_docx_sources(sample)

    assert [source.name for source in sources] == [
        "欧盟安全事务研究所",
        "欧洲政策研究中心",
        "马腾斯欧洲事务研究中心",
        "艾格蒙特国际关系研究所",
        "欧洲政策中心",
        "西班牙埃尔卡诺皇家研究所",
    ]
    assert sources[0].adapter_key == "euiss"
    assert str(sources[1].entries[0].url) == "https://www.ceps.eu/"


def test_parses_heading_inline_urls_and_table_rows() -> None:
    document = Document()
    document.add_heading("欧洲智库", level=1)
    document.add_paragraph("1. Example Institute")
    document.add_paragraph("https://example.com/")
    document.add_paragraph("2. Inline Centre https://inline.example.org/research")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table Institute"
    table.cell(0, 1).text = "https://table.example.net/publications"
    buffer = io.BytesIO()
    document.save(buffer)

    sources = parse_docx_sources(buffer.getvalue())

    assert len(sources) == 3
    assert all(source.group_name == "欧洲智库" for source in sources)
    assert sources[1].name == "Inline Centre"
    assert sources[2].name == "Table Institute"


def test_deduplicates_urls() -> None:
    document = Document()
    document.add_paragraph("One")
    document.add_paragraph("https://example.com/?utm_source=test")
    document.add_paragraph("Duplicate")
    document.add_paragraph("https://example.com/")
    buffer = io.BytesIO()
    document.save(buffer)

    assert len(parse_docx_sources(buffer.getvalue())) == 1


def test_sources_json_preserves_enabled_and_rejects_future_schema() -> None:
    rows = parse_sources_json(
        '{"schema_version":1,"sources":[{"name":"One","homepage_url":"https://example.org","enabled":false}]}'
    )
    assert rows[0].enabled is False
    assert str(rows[0].entries[0].url) == "https://example.org/"
    v2_rows = parse_sources_json(
        '{"schema_version":2,"sources":[{"name":"One","homepage_url":"https://example.org/",'
        '"entries":[{"url":"https://example.org/research/","adapter_key":"generic"},'
        '{"url":"https://example.org/reports/","adapter_key":"generic"}]}]}'
    )
    assert [str(entry.url) for entry in v2_rows[0].entries] == [
        "https://example.org/research/",
        "https://example.org/reports/",
    ]
    try:
        parse_sources_json('{"schema_version":99,"sources":[]}')
    except ValueError as exc:
        assert "更高版本" in str(exc)
    else:
        raise AssertionError("future schema must be rejected")


def test_groups_multiple_urls_under_one_known_institute_without_replacing_them() -> None:
    document = Document()
    document.add_paragraph("Elcano Institute")
    document.add_paragraph("https://www.realinstitutoelcano.org/en/publications/")
    document.add_paragraph("https://www.realinstitutoelcano.org/en/reports/")
    buffer = io.BytesIO()
    document.save(buffer)

    sources = parse_docx_sources(buffer.getvalue())

    assert len(sources) == 1
    assert [str(entry.url) for entry in sources[0].entries] == [
        "https://www.realinstitutoelcano.org/en/publications/",
        "https://www.realinstitutoelcano.org/en/reports/",
    ]
    assert all(entry.adapter_key == "elcano" for entry in sources[0].entries)


def test_generated_word_example_round_trips_names_groups_and_multiple_entries(tmp_path) -> None:
    path = create_word_example(tmp_path / "sources-example.docx")
    sources = parse_docx_sources(path)

    assert [(source.name, source.group_name) for source in sources] == [
        ("示例智库甲", "欧洲与国际战略"),
        ("示例智库乙", "国际战略"),
    ]
    assert [str(entry.url) for entry in sources[0].entries] == [
        "https://example.org/publications/",
        "https://example.org/reports/",
    ]
    assert str(sources[1].entries[0].url) == "https://research.example.net/papers/"
