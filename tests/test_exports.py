from __future__ import annotations

import json
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from sqlalchemy import select

from thinktank_tracker.constants import CandidateStatus, Decision
from thinktank_tracker.models import Candidate, CandidateAttachment, CrawlRun, ReviewDecision, Source, SourceEntry, SourceEntryRun, SourceRun
from thinktank_tracker.paths import PATHS
import thinktank_tracker.exporters as exporters


def _seed(db_session):
    source = Source(name="示例智库", homepage_url="https://example.org/", publication_url="https://example.org/publications")
    empty_source = Source(name="无符合项智库", homepage_url="https://empty.example/", publication_url="https://empty.example/research")
    incomplete_source = Source(name="覆盖不完整智库", homepage_url="https://partial.example/", publication_url="https://partial.example/research")
    run = CrawlRun(start_date=date(2026, 8, 2), end_date=date(2026, 8, 26), status="completed")
    db_session.add_all([source, empty_source, incomplete_source, run])
    db_session.flush()
    entries = [
        SourceEntry(source_id=source.id, url=source.publication_url, canonical_url=source.publication_url),
        SourceEntry(source_id=empty_source.id, url=empty_source.publication_url, canonical_url=empty_source.publication_url),
        SourceEntry(source_id=incomplete_source.id, url=incomplete_source.publication_url, canonical_url=incomplete_source.publication_url),
    ]
    db_session.add_all(entries)
    db_session.flush()
    source_runs = [
        SourceRun(run_id=run.id, source_id=source.id, status="complete", pages_scanned=2, links_discovered=1, stop_reason="已到栏目末页"),
        SourceRun(run_id=run.id, source_id=empty_source.id, status="complete", pages_scanned=1, links_discovered=0, stop_reason="已到栏目末页"),
        SourceRun(run_id=run.id, source_id=incomplete_source.id, status="incomplete", pages_scanned=25, links_discovered=0, stop_reason="达到安全上限"),
    ]
    db_session.add_all(source_runs)
    db_session.flush()
    db_session.add_all([
        SourceEntryRun(run_id=run.id, source_run_id=row.id, source_entry_id=entry.id, status=row.status, pages_scanned=row.pages_scanned, links_discovered=row.links_discovered, stop_reason=row.stop_reason)
        for row, entry in zip(source_runs, entries, strict=True)
    ])
    item = Candidate(
        run_id=run.id, source_id=source.id, canonical_url="https://example.org/report",
        original_url="https://example.org/report", title_original="Strategic Competition",
        title_en="Strategic Competition", title_zh="战略竞争", published_date=date(2026, 8, 26),
        date_status="confirmed", date_evidence="meta", content_type="report", language="en",
        english_word_count=2200, extraction_status="complete", content_hash="a" * 64,
        extracted_text="full text", status=CandidateStatus.APPROVED,
    )
    db_session.add(item)
    db_session.flush()
    db_session.add(
        CandidateAttachment(
            candidate_id=item.id,
            url="https://example.org/report.pdf",
            canonical_url="https://example.org/report.pdf",
            role="primary",
            link_text="Download report",
            extraction_status="complete",
            used_as_body=True,
            extraction_audit_json=json.dumps(
                {
                    "final_engine": "pypdf",
                    "final_engine_version": "6.16.2",
                    "duration_ms": 12.5,
                    "fallback_reason": "anydoc_malformed",
                    "fallback_detail": "AnyDoc 转换失败: broken xref",
                    "ocr_pages": [2, 4],
                    "final_source": "pypdf_text",
                    "attempts": [],
                },
                ensure_ascii=False,
            ),
        )
    )
    db_session.add(ReviewDecision(
        candidate_id=item.id, decision=Decision.INCLUDE, primary_category="macro",
        secondary_categories_json='["technology_competition"]', title_en="Strategic Competition",
        title_zh="战略竞争", authors_json='["Alice Example"]', content_type="report",
        core_claim_zh="竞争正在重塑国际秩序。",
    ))
    db_session.commit()
    return run


def test_docx_and_xlsx_exports_have_required_structure(db_session, tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(exporters, "PATHS", replace(PATHS, exports=tmp_path))
    run = _seed(db_session)
    docx_path = exporters.export_report(db_session, run.id)
    xlsx_path = exporters.export_audit_workbook(db_session, run.id)
    assert docx_path.exists() and xlsx_path.exists()
    assert docx_path.name == "2026年8月2日-2026年8月26日.docx"
    document = Document(docx_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.tables == []
    assert document.paragraphs[0].text == "智库周览（2026年8月2日-2026年8月26日）"
    assert document.paragraphs[0].alignment == 1
    assert document.paragraphs[1].text == "示例智库"
    assert document.paragraphs[2].text == (
        "Alice Example, Strategic Competition（战略竞争）, August 26, 2026."
    )
    assert not document.paragraphs[2].text.startswith("1.")
    assert [run.italic for run in document.paragraphs[2].runs] == [None, True, None]
    assert "无符合项智库" not in text
    assert "覆盖不完整智库" not in text
    assert document.paragraphs[3]._p.xpath(".//w:hyperlink")
    assert document.paragraphs[3].paragraph_format.keep_together is True
    hyperlink_text = "".join(document.paragraphs[3]._p.xpath(".//w:t/text()"))
    assert hyperlink_text == "https://example.org/report"
    relationship_id = document.paragraphs[3]._p.xpath(".//w:hyperlink")[0].get(qn("r:id"))
    assert document.paragraphs[3].part.rels[relationship_id].target_ref == "https://example.org/report"
    section = document.sections[0]
    assert section.page_width.inches == pytest.approx(8.2677, abs=0.01)
    assert section.page_height.inches == pytest.approx(11.6929, abs=0.01)
    assert section.top_margin.inches == pytest.approx(1.0, abs=0.01)
    assert section.bottom_margin.inches == pytest.approx(1.0, abs=0.01)
    assert section.left_margin.inches == pytest.approx(1.25, abs=0.01)
    assert section.right_margin.inches == pytest.approx(1.25, abs=0.01)
    workbook = load_workbook(xlsx_path, read_only=True)
    assert workbook.sheetnames == [
        "概览", "智库覆盖", "入口覆盖", "逐页发现审计", "无符合项智库", "日期待确认", "抓取失败",
        "全部正式内容", "日期范围内完整清单", "范围外内容", "模型调用统计", "PDF附件", "人工触发模型建议", "模型日期辅助", "研究判定核验",
    ]
    assert workbook["全部正式内容"].max_row == 2
    attachment_rows = list(workbook["PDF附件"].iter_rows(values_only=True))
    assert attachment_rows[0][7:13] == (
        "解析器",
        "解析器版本",
        "解析耗时（毫秒）",
        "回退原因",
        "OCR页码",
        "最终正文来源",
    )
    assert attachment_rows[1][7:13] == (
        "pypdf",
        "6.16.2",
        12.5,
        "AnyDoc 转换失败: broken xref",
        "2, 4",
        "pypdf_text",
    )


def test_exports_use_task_source_and_entry_snapshots(db_session, tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(exporters, "PATHS", replace(PATHS, exports=tmp_path))
    run = _seed(db_session)
    source_run = db_session.scalar(
        select(SourceRun)
        .where(SourceRun.run_id == run.id)
        .order_by(SourceRun.id)
    )
    entry_run = source_run.entry_runs[0]
    source_run.source_name_snapshot = "任务创建时智库"
    entry_run.entry_url_snapshot = "https://snapshot.example/reports/"
    source_run.source.name = "后来改名的智库"
    entry_run.source_entry.url = "https://changed.example/reports/"
    db_session.commit()

    docx_path = exporters.export_report(db_session, run.id)
    xlsx_path = exporters.export_audit_workbook(db_session, run.id)
    document = Document(docx_path)
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "任务创建时智库" in document_text
    assert "后来改名的智库" not in document_text

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    assert workbook["全部正式内容"]["B2"].value == "任务创建时智库"
    assert workbook["入口覆盖"]["B2"].value == "https://snapshot.example/reports/"


def test_report_export_fails_as_one_batch_when_an_approved_item_has_no_author(
    db_session, tmp_path, monkeypatch
) -> None:
    monkeypatch.setattr(exporters, "PATHS", replace(PATHS, exports=tmp_path))
    run = _seed(db_session)
    item = db_session.scalar(select(Candidate).where(Candidate.run_id == run.id))
    item.review.authors_json = "[]"
    db_session.commit()

    with pytest.raises(ValueError, match=rf"候选 {item.id}: 作者"):
        exporters.export_report(db_session, run.id)
    assert list(tmp_path.glob("*.docx")) == []


def test_report_orders_sources_and_numbers_only_multi_item_groups(
    db_session, tmp_path, monkeypatch
) -> None:
    monkeypatch.setattr(exporters, "PATHS", replace(PATHS, exports=tmp_path))
    run = _seed(db_session)
    first_source, second_source = db_session.scalars(
        select(Source).order_by(Source.id).limit(2)
    ).all()

    earlier = Candidate(
        run_id=run.id,
        source_id=first_source.id,
        canonical_url="https://example.org/earlier",
        original_url="https://example.org/earlier",
        title_original="Earlier",
        published_date=date(2026, 8, 7),
        date_status="confirmed",
        content_type="policy brief",
        status=CandidateStatus.APPROVED,
    )
    second_group = Candidate(
        run_id=run.id,
        source_id=second_source.id,
        canonical_url="https://empty.example/only",
        original_url="https://empty.example/only",
        title_original="Only",
        published_date=date(2026, 8, 4),
        date_status="confirmed",
        content_type="report",
        status=CandidateStatus.APPROVED,
    )
    db_session.add_all([earlier, second_group])
    db_session.flush()
    db_session.add_all(
        [
            ReviewDecision(
                candidate_id=earlier.id,
                decision=Decision.INCLUDE,
                primary_category="macro",
                title_en="A much earlier title",
                title_zh="更早的标题",
                authors_json='["Earlier Author"]',
                content_type="policy brief",
                core_claim_zh="观点",
            ),
            ReviewDecision(
                candidate_id=second_group.id,
                decision=Decision.INCLUDE,
                primary_category="macro",
                title_en="Only title",
                title_zh="唯一标题",
                authors_json='["Only Author"]',
                content_type="report",
                core_claim_zh="观点",
            ),
        ]
    )
    db_session.commit()

    document = Document(exporters.export_report(db_session, run.id))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs[1] == "示例智库"
    assert paragraphs[2].startswith("1. Earlier Author")
    assert paragraphs[4].startswith("2. Alice Example")
    second_heading = paragraphs.index("无符合项智库")
    assert paragraphs[second_heading + 1].startswith("Only Author")
    assert not paragraphs[second_heading + 1].startswith("1.")


def test_internal_template_is_blank_and_user_example_is_complete() -> None:
    project_root = Path(__file__).resolve().parents[1]
    template = Document(
        project_root / "app" / "thinktank_tracker" / "assets" / "report_template.docx"
    )
    assert template.paragraphs == []
    assert template.tables == []

    example = Document(project_root / "docs" / "示例文件" / "智库周览格式示例.docx")
    paragraphs = example.paragraphs
    paragraph_text = [paragraph.text for paragraph in paragraphs]
    assert example.tables == []
    assert paragraph_text[0] == "智库周览（2026年8月7日-2026年8月21日）"
    assert paragraph_text[1] == "示例智库甲（Example Institute A）"
    assert not paragraph_text[2].startswith("1. ")
    assert paragraph_text[4] == "示例智库乙（Example Institute B）"
    assert paragraph_text[5].startswith("1. Bob Example, Carol Example")
    assert paragraph_text[7].startswith("2. David Example")
    assert "Technology Competition, Critical Infrastructure" in paragraph_text[7]
    assert paragraph_text[9] == "示例智库丙（Example Institute C）"
    assert paragraph_text[10].startswith("1. Elena Example, Farid Example")
    assert paragraph_text[-2].startswith("8. Elena Example, Farid Example")

    hyperlink_paragraphs = [
        paragraph for paragraph in paragraphs if paragraph._p.xpath(".//w:hyperlink")
    ]
    assert len(hyperlink_paragraphs) == 11
    assert all(paragraph.paragraph_format.keep_together is True for paragraph in hyperlink_paragraphs)
    assert all(
        paragraph.part.rels[
            paragraph._p.xpath(".//w:hyperlink")[0].get(qn("r:id"))
        ].is_external
        for paragraph in hyperlink_paragraphs
    )

    article_runs = paragraphs[5].runs
    assert [run.italic for run in article_runs] == [None, None, True, None]
    assert paragraphs[0].runs[0].bold is True
    assert paragraphs[1].runs[0].bold is True

    section = example.sections[0]
    assert section.page_width.inches == pytest.approx(8.2677, abs=0.01)
    assert section.page_height.inches == pytest.approx(11.6929, abs=0.01)
    assert section.top_margin.inches == pytest.approx(1.0, abs=0.01)
    assert section.bottom_margin.inches == pytest.approx(1.0, abs=0.01)
    assert section.left_margin.inches == pytest.approx(1.25, abs=0.01)
    assert section.right_margin.inches == pytest.approx(1.25, abs=0.01)
